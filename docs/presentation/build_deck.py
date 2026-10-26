from pathlib import Path
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image,ImageDraw,ImageFont
from reportlab.pdfgen import canvas
import textwrap
D=Path(__file__).parent
prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
W,H=1600,900
BG='#0B1820'; PANEL='#142731';WHITE='#F1F6F7';MUTED='#ABC1CA';TEAL='#56E3CF';LINE='#2C4651'
font='/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf';bold='/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf'
slides=[];notes=[]
def rgb(c):return RGBColor.from_string(c.strip('#'))
def box(x,y,w,h,c,r=0):
 sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if r else MSO_SHAPE.RECTANGLE, Inches(x/120), Inches(y/120), Inches(w/120), Inches(h/120));sh.fill.solid();sh.fill.fore_color.rgb=rgb(c);sh.line.fill.background()
 if r: draw.rounded_rectangle((x,y,x+w,y+h),radius=r,fill=c)
 else:draw.rectangle((x,y,x+w,y+h),fill=c)
def text(t,x,y,w,size=32,color=WHITE,weight=False):
 f=ImageFont.truetype(bold if weight else font,size)
 lines=[]
 for para in t.split('\n'):
  line=''
  for word in para.split():
   q=(line+' '+word).strip()
   if draw.textlength(q,font=f)>w and line:lines.append(line);line=word
   else:line=q
  lines.append(line)
 rendered='\n'.join(lines)
 draw.multiline_text((x,y),rendered,font=f,fill=color,spacing=int(size*.3))
 tb=s.shapes.add_textbox(Inches(x/120),Inches(y/120),Inches(w/120),Inches((len(lines)*size*1.35+20)/120));tf=tb.text_frame;tf.word_wrap=False;tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
 for i,l in enumerate(lines):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=l;p.font.name='DejaVu Sans';p.font.size=Pt(size*.6);p.font.bold=weight;p.font.color.rgb=rgb(color);p.space_after=Pt(size*.18)
 return len(lines)*size*1.3

def picture(name,x,y,w,h):
 im=Image.open(D/name).convert('RGB');ratio=min(w/im.width,h/im.height);iw=int(im.width*ratio);ih=int(im.height*ratio);xx=x+(w-iw)//2;yy=y+(h-ih)//2
 box(xx-10,yy-10,iw+20,ih+20,LINE,22)
 img.paste(im.resize((iw,ih),Image.Resampling.LANCZOS),(xx,yy))
 s.shapes.add_picture(str(D/name), Inches(xx/120), Inches(yy/120),width=Inches(iw/120),height=Inches(ih/120))
def start(k,title=None):
 global s,img,draw
 s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=rgb(BG)
 img=Image.new('RGB',(W,H),BG);draw=ImageDraw.Draw(img)
 text('ReachAI',65,35,350,25,TEAL,True);text(k.upper(),870,39,660,17,MUTED)
 box(65,91,1470,2,LINE)
 if title:text(title,65,135,1435,54,WHITE,True)
 text('CONGRESSIONAL APP CHALLENGE',65,856,1000,15,MUTED);text(f'{len(slides)+1:02}',1460,850,65,22,TEAL)
def end(note):
 path=D/f'slide-{len(slides)+1:02}.png';img.save(path);slides.append(path);notes.append(note);s.notes_slide.notes_text_frame.text=note

start('A practical next step')
text('Find food.\nFind a way forward.',65,225,980,78,WHITE,True)
text('Food assistance, made easier to find.',70,465,900,34,MUTED)
box(70,580,600,66,TEAL,16);text('ReachAI  /  Congressional App Challenge',93,598,560,22,BG,True)
text('Raman Jamal · Bryan High · 11th grade\nAran Jamal · Davila · 7th grade',70,700,960,25,MUTED)
picture('home.png',1110,125,380,680)
end('Raman: Hi, I am Raman Jamal, an 11th-grade student at Bryan High. Aran: And I am Aran Jamal, a 7th-grade student at Davila. Together: This is ReachAI, designed to make finding food assistance simpler. This deck shows actual local app screenshots captured October 25, 2026 Central time.')

start('The problem','A national problem. A local next step.')
text('≈ 1 in 7',70,240,900,132,TEAL,True)
text('U.S. households experienced\nfood insecurity in 2024.',80,420,850,42,WHITE,True)
for i in range(7):box(85+i*111,565,82,82,TEAL if i==0 else PANEL,18)
box(1040,250,485,420,PANEL,24);text('47.9M',1080,300,400,88,TEAL,True);text('people lived in\nfood-insecure households',1080,430,390,29)
text('USDA ERS · 2024 data · 13.7% of households ≈ one in seven',80,735,1430,22,MUTED)
end('USDA reports that 13.7 percent of U.S. households experienced food insecurity in 2024—about one in seven. 47.9 million people lived in food-insecure households. This is not a pantry-use statistic and does not measure lack of awareness. Source: https://www.ers.usda.gov/topics/food-nutrition-assistance/food-security-in-the-us/key-statistics-graphics')

start('The design question','Make the next decision easier.')
for i,(n,t,d) in enumerate([('01','Where can I go?','Source-backed pickup locations.'),('02','Will it fit my needs?','Pickup times and known requirements.'),('03','How do I get there?','Location-aware directions and contact details.')]):
 x=65+i*500;box(x,285,470,380,PANEL,22);text(n,x+30,320,400,40,TEAL,True);text(t,x+30,400,400,36,WHITE,True);text(d,x+30,535,390,27,MUTED)
end('ReachAI focuses on three practical questions: where can I go, will it fit my needs, and how do I get there? The design prioritizes simple actions and readable information over long forms. We have not measured how many people miss pantry help because of lack of information.')

start('Demo 01 / Search','Ask in everyday language.')
text('“I need food in Bryan\nand I’m walking.”',70,290,900,55,TEAL,True)
text('No account needed to search.\nUse a location or enter an address.\nKeep browsing if AI is unavailable.',75,510,920,31,MUTED)
picture('home.png',1100,210,370,600)
end('On the home screen, describe what you need. A user can choose a starting location or enter an address. Food search does not require an account. The current screenshots are real local captures. The AI checkbox is a preference, not evidence that a live AI request succeeded.')

start('Demo 02 / Results','Useful details, not just a pin.')
text('66',70,250,800,130,TEAL,True);text('maintained food-assistance sites',75,410,900,36,WHITE,True)
text('Pickup address and published schedule\nKnown requirements and contact details\nSource information and uncertainty',75,520,880,29,MUTED)
text('Includes restricted programs. Coverage is not exhaustive.',75,745,880,22,MUTED)
picture('result.png',1070,205,430,610)
end('The maintained directory has 66 pantry and food-distribution sites after source and district checks. They are not all unrestricted walk-in pantries: some are senior grocery programs. This screenshot is a directory result, not a live AI recommendation. Food stock and pickup entrances have not been verified on-site.')

start('AI + reliability','Help explain. Never invent.')
for x,t,d in [(65,'Compare real options','Distance, published pickup windows, and known requirements inform comparisons.'),(565,'Explain the choices','When configured, hosted AI helps interpret requests and summarize listed facts.'),(1065,'Keep a fallback','Directory results remain available when AI cannot respond.')]:
 box(x,275,470,415,PANEL,22);text(t,x+30,320,405,37,TEAL,True);text(d,x+30,455,400,29)
text('No fabricated locations, live stock claims, or guaranteed “perfect” match.',75,755,1430,25,MUTED)
end('The app combines comparison logic with AI-assisted interpretation and explanation. AI must use the directory facts. This local session had no configured Groq key, so live AI was not exercised for these slides. Do not claim a successful live summary unless you demonstrate one during recording. The app cannot guarantee the perfect pantry.')

start('Demo 03 / Community','Share food. Review before publishing.')
text('Host a Food Pickup',70,280,900,45,TEAL,True);text('Community meals and pickup events.',75,350,870,30,MUTED)
text('Share Extra Food',70,450,900,45,TEAL,True);text('Sealed, shelf-stable groceries.',75,520,870,30,MUTED)
text('Submit → Moderator review → Public listing',75,690,925,29,WHITE,True)
picture('share.png',1100,205,370,610)
end('Community members can submit pickup events or sealed, shelf-stable groceries without creating an account. Listings require moderator approval before they become public. Approved active listings can appear alongside pantry options. No real submission was created while capturing this deck. Review is not a guarantee of food safety or availability.')

start('Engineering','Built for a dependable core.')
items=[('INTERFACE','Next.js + React','Mobile-first screens'),('APPLICATION','TypeScript','Search and comparison logic'),('STORAGE','SQLite','Listings and moderation'),('EXTERNAL SERVICES','Maps + hosted AI','Honest failure states')]
for i,(a,b,c) in enumerate(items):
 x=65+(i%2)*750;y=260+(i//2)*250;box(x,y,715,215,PANEL,20);text(a,x+30,y+25,650,18,TEAL,True);text(b,x+30,y+68,650,38,WHITE,True);text(c,x+30,y+135,650,27,MUTED)
end('ReachAI uses Next.js, React, TypeScript, and SQLite. OpenStreetMap-based mapping and hosted AI are external services. Search and source-backed information remain useful when those services fail. Explain your own contributions honestly and disclose AI development assistance where submission rules require it.')

start('Evidence / limits','Tested software. More fieldwork ahead.')
box(65,265,720,450,PANEL,22);text('661',100,300,630,105,TEAL,True);text('unit tests passed',105,430,620,35,WHITE,True);text('Latest UI check: 14 mobile / desktop checks\nTypeScript + production build passed',105,535,620,27,MUTED)
text('Next priorities',870,280,640,43,WHITE,True);text('Test with real users.\nConfirm pickup details with providers.\nVerify the deployed demo end to end.',875,385,640,32,MUTED)
text('Tests do not prove food availability or real-world impact.',875,630,620,25,TEAL)
end('The recorded project checks include 661 unit tests and 14 mobile/desktop browser checks for the latest UI changes; TypeScript and production build passed. This is not a new full audit of every workflow. During this screenshot session AI was unconfigured and the basemap failed to load, so no successful live AI or map demonstration is claimed. Next priorities are real-user testing, provider confirmation and deployed installation verification.')

start('The goal')
text('From “I need food”\nto a practical next step.',75,245,1450,83,WHITE,True)
text('ReachAI',80,535,1400,76,TEAL,True);text('Raman Jamal + Aran Jamal',85,675,1350,36,WHITE,True);text('Thank you',85,745,1350,28,MUTED)
end('ReachAI does not solve food insecurity alone. Its goal is to make a practical next step easier. We are Raman Jamal and Aran Jamal, and this is ReachAI. Thank you. Add a real public demo URL or QR code only after deployment is verified; none has been invented here.')

start('Appendix / sources','Evidence behind the presentation')
refs=[('USDA Economic Research Service','2024: 13.7% of U.S. households; 47.9 million people in food-insecure households.','https://www.ers.usda.gov/topics/food-nutrition-assistance/food-security-in-the-us/key-statistics-graphics'),('Directory source review','66 maintained sites; restrictions and unresolved facts retained.','Repository: docs/TX10-COVERAGE-AUDIT.md and docs/TX10-SIXTY-REVIEW.md'),('Real app screenshots','Captured from the local ReachAI build. No invented AI response or private submission.','Repository: https://github.com/returnedlegend/Reach  •  branch improve/food-first-experience')]
for i,(a,b,c) in enumerate(refs):
 y=250+i*175;text(a,75,y,1450,30,TEAL,True);text(b,75,y+48,1440,24);text(c,75,y+92,1440,17,MUTED)
end('Source appendix; keep available for judges, but it need not be narrated. Source statistics were directly checked on the USDA webpage in this conversation. Screenshot capture: October25,2026 America/Chicago. Software checks are historical reported checks from the latest changes, not user impact measurements. Presenters: Raman Jamal, Bryan High, 11th grade; Aran Jamal, Davila, 7th grade.')

prs.save(D/'ReachAI-Congressional-App-Challenge.pptx')
c=canvas.Canvas(str(D/'ReachAI-Congressional-App-Challenge.pdf'),pagesize=(960,540))
for p in slides:c.drawImage(str(p),0,0,width=960,height=540);c.showPage()
c.save()
(D/'Speaker-notes.txt').write_text('\n\n'.join(f'SLIDE {i+1}\n{n}' for i,n in enumerate(notes)))
thumb=Image.new('RGB',(1200,675*4),BG)
for i,p in enumerate(slides):thumb.paste(Image.open(p).resize((600,338)),((i%2)*600,(i//2)*450))
thumb.save(D/'slide-overview.jpg')
print('Created',len(slides),'slides: editable PowerPoint, PDF, notes and preview.')
